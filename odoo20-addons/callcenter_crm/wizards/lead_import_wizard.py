import base64
import csv
import hashlib
import io
import re

from openpyxl import load_workbook

from odoo import fields, models, _
from odoo.exceptions import AccessError, UserError, ValidationError


_HEADER_RE = re.compile(r"[^a-z0-9]+")
_HEADER_MAP = {
    "contactname": "contact_name",
    "name": "contact_name",
    "phone": "phone",
    "telephone": "phone",
    "mobile": "phone",
    "email": "email",
    "externalid": "external_id",
    "company": "company",
    "priority": "priority",
    "source": "source",
    "notes": "notes",
}
_FORBIDDEN = {
    "campaign", "supervisor", "assignedagent", "agent", "queuestate",
    "createdby", "createddate", "userid", "teamid",
}


class CallCenterLeadImportWizard(models.TransientModel):
    _name = "callcenter.lead.import.wizard"
    _description = "Campaign Lead Import"

    campaign_id = fields.Many2one("callcenter.campaign", required=True, ondelete="cascade")
    upload_file = fields.Binary(required=True, attachment=False)
    filename = fields.Char(required=True)
    allow_duplicate_file = fields.Boolean(
        string="Deliberately Re-import Duplicate File",
        help="Use only after reviewing the warning for a previously imported identical file.",
    )
    batch_id = fields.Many2one("callcenter.lead.import.batch", readonly=True)

    def _ensure_admin(self):
        if not (
            self.env.is_superuser()
            or self.env.user.has_group("base.group_system")
            or self.env.user.has_group("callcenter_crm.group_callcenter_superuser")
        ):
            raise AccessError(_("Only the Call Center Super User may import call-center lead lists."))

    def _normalized_header(self, value):
        return _HEADER_RE.sub("", str(value or "").strip().lower())

    def _canonicalize_rows(self, headers, rows):
        normalized = [self._normalized_header(h) for h in headers]
        forbidden = sorted(set(normalized) & _FORBIDDEN)
        if forbidden:
            raise ValidationError(
                _("The spreadsheet may not control protected fields: %s", ", ".join(forbidden))
            )
        mapping = [_HEADER_MAP.get(h) for h in normalized]
        if "contact_name" not in mapping or "phone" not in mapping:
            raise ValidationError(_("The import requires Contact Name and Phone columns."))
        result = []
        for values in rows:
            row = {}
            for index, canonical in enumerate(mapping):
                if not canonical:
                    continue
                value = values[index] if index < len(values) else None
                if value is None:
                    value = ""
                elif not isinstance(value, (str, int, float, bool)):
                    value = str(value)
                row[canonical] = str(value).strip()
            result.append(row)
        return result

    def _read_rows(self, raw):
        lower = (self.filename or "").lower()
        if lower.endswith(".csv"):
            text = raw.decode("utf-8-sig")
            reader = csv.reader(io.StringIO(text))
            all_rows = list(reader)
            if not all_rows:
                return []
            return self._canonicalize_rows(all_rows[0], all_rows[1:])
        if lower.endswith(".xlsx"):
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            sheet = workbook.active
            values = list(sheet.iter_rows(values_only=True))
            if not values:
                return []
            return self._canonicalize_rows(values[0], values[1:])
        raise ValidationError(_("Only CSV and XLSX lead lists are supported."))

    def _priority(self, value):
        value = (value or "").strip().lower()
        return {
            "": "0", "low": "0", "0": "0",
            "medium": "1", "normal": "1", "1": "1",
            "high": "2", "2": "2",
            "very high": "3", "veryhigh": "3", "3": "3",
        }.get(value, "0")

    def action_validate(self):
        self.ensure_one()
        self._ensure_admin()
        raw = base64.b64decode(self.upload_file or b"")
        digest = hashlib.sha256(raw).hexdigest()
        previous = self.env["callcenter.lead.import.batch"].sudo().search([
            ("campaign_id", "=", self.campaign_id.id),
            ("file_hash", "=", digest),
            ("state", "=", "imported"),
        ], limit=1)
        if previous and not self.allow_duplicate_file:
            raise UserError(
                _("This exact file was already imported into this campaign. Review the prior batch and enable the deliberate override to continue.")
            )

        rows = self._read_rows(raw)
        batch = self.env["callcenter.lead.import.batch"].sudo().create({
            "campaign_id": self.campaign_id.id,
            "filename": self.filename,
            "file_hash": digest,
            "imported_by_id": self.env.user.id,
            "imported_at": fields.Datetime.now(),
            "state": "draft",
            "total_rows": len(rows),
        })
        Lead = self.env["crm.lead"].sudo().with_context(active_test=False)
        seen = set()
        duplicate_count = 0
        error_count = 0
        line_values = []

        for number, row in enumerate(rows, start=2):
            phone_key = Lead._cc_phone_normalize(row.get("phone"))
            email_key = Lead._cc_email_normalize(row.get("email"))
            external_id = (row.get("external_id") or "").strip() or False
            status = "ready"
            message = False
            if not row.get("contact_name") or not phone_key:
                status = "error"
                message = _("Contact Name and Phone are required.")
                error_count += 1
            else:
                logical_key = ("phone", phone_key)
                if external_id:
                    external_key = ("external", external_id)
                else:
                    external_key = None
                if logical_key in seen or (external_key and external_key in seen):
                    status = "duplicate"
                    message = _("Duplicate row inside the uploaded file.")
                    duplicate_count += 1
                else:
                    duplicate = Lead._cc_find_duplicate(
                        self.campaign_id.id, phone_key, email_key, external_id
                    )
                    if duplicate:
                        status = "duplicate"
                        message = _("Duplicate of existing lead %s", duplicate.display_name)
                        duplicate_count += 1
                    seen.add(logical_key)
                    if external_key:
                        seen.add(external_key)
            line_values.append({
                "batch_id": batch.id,
                "row_number": number,
                "raw_data": row,
                "normalized_phone": phone_key,
                "normalized_email": email_key,
                "status": status,
                "error_message": message or False,
            })

        self.env["callcenter.lead.import.line"].sudo().create(line_values)
        batch.sudo().write({
            "state": "validated",
            "duplicate_count": duplicate_count,
            "error_count": error_count,
        })
        self.batch_id = batch
        return {
            "type": "ir.actions.act_window",
            "name": _("Validated Lead Import"),
            "res_model": "callcenter.lead.import.batch",
            "res_id": batch.id,
            "view_mode": "form",
            "target": "current",
        }

    def action_import(self):
        self.ensure_one()
        self._ensure_admin()
        if not self.batch_id:
            self.action_validate()
        batch = self.batch_id.sudo()
        if batch.state != "validated":
            raise ValidationError(_("Only a validated batch can be imported."))

        created = 0
        errors = batch.error_count
        Lead = self.env["crm.lead"].sudo()
        for line in batch.line_ids.filtered(lambda l: l.status == "ready"):
            row = line.raw_data
            vals = {
                "name": row.get("contact_name"),
                "contact_name": row.get("contact_name"),
                "phone": row.get("phone"),
                "email_from": row.get("email") or False,
                "partner_name": row.get("company") or False,
                "priority": self._priority(row.get("priority")),
                "description": row.get("notes") or False,
                "cc_source_text": row.get("source") or False,
                "source_external_id": row.get("external_id") or False,
                "cc_campaign_id": batch.campaign_id.id,
                "source_batch_id": batch.id,
                "user_id": False,
                "queue_state": "available",
            }
            try:
                with self.env.cr.savepoint():
                    lead = Lead.create(vals)
                    line.sudo().write({"status": "created", "lead_id": lead.id})
                    created += 1
            except Exception as exc:
                line.sudo().write({"status": "error", "error_message": str(exc)[:500]})
                errors += 1
        batch.sudo().write({
            "created_count": created,
            "error_count": errors,
            "state": "imported" if created or not errors else "failed",
        })
        return {
            "type": "ir.actions.act_window",
            "name": _("Lead Import Batch"),
            "res_model": "callcenter.lead.import.batch",
            "res_id": batch.id,
            "view_mode": "form",
            "target": "current",
        }
