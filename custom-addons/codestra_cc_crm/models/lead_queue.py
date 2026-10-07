import base64
import csv
import hashlib
import io
import re

from odoo import _, api, fields, models
from odoo.exceptions import AccessError, UserError, ValidationError

from .crm_workspace import (
    CRM_TRANSITION_CAPABILITY,
    _is_global_admin,
    _is_operational,
)


QUEUE_SYSTEM_CAPABILITY = object()
ASSIGNMENT_HISTORY_CAPABILITY = object()
QUEUE_CONTROLLED_FIELDS = {
    "user_id",
    "queue_state",
    "next_eligible_at",
    "call_attempt_count",
    "last_call_at",
}
TRANSITION_CONTROLLED_FIELDS = {
    "codestra_workflow_id",
    "codestra_current_status_id",
    "codestra_previous_status_id",
    "status_entered_at",
    "stage_id",
}
QUEUE_ACTIVE_LIFECYCLE_STATES = {"staging_ready", "active"}
NON_DIGIT_RE = re.compile(r"\D+")
COLUMN_ALIASES = {
    "name": "name",
    "contact_name": "name",
    "contact": "name",
    "phone": "phone",
    "telephone": "phone",
    "mobile": "phone",
    "external_lead_id": "external_id",
    "external_id": "external_id",
    "lead_id": "external_id",
    "email": "email",
    "company": "company",
    "company_name": "company",
    "priority": "priority",
    "source": "source",
    "notes": "notes",
    "note": "notes",
}


def _phone_key(value):
    digits = NON_DIGIT_RE.sub("", str(value or ""))
    return digits or False


def _email_key(value):
    value = str(value or "").strip().lower()
    return value or False


def _canonical_column(value):
    value = str(value or "").strip().lower()
    value = re.sub(r"[^a-z0-9]+", "_", value).strip("_")
    return COLUMN_ALIASES.get(value, value)


def _json_value(value):
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _priority_value(value):
    raw = str(value or "").strip().lower()
    aliases = {
        "very low": "0",
        "low": "1",
        "normal": "1",
        "medium": "2",
        "high": "3",
        "very high": "3",
        "urgent": "3",
    }
    if raw in aliases:
        return aliases[raw]
    if raw.isdigit():
        return str(max(0, min(3, int(raw))))
    return "0"


def _require_global_admin(env):
    if not _is_global_admin(env.user):
        raise AccessError(_("Only the Call Center Super User may perform lead imports."))


class CrmLeadQueue(models.Model):
    _inherit = "crm.lead"

    source_batch_id = fields.Many2one(
        "callcenter.lead.import.batch",
        ondelete="restrict",
        index=True,
        copy=False,
        readonly=True,
    )
    source_external_id = fields.Char(index=True, copy=False)
    queue_state = fields.Selection(
        [
            ("available", "Available"),
            ("assigned", "Assigned"),
            ("callback", "Callback"),
            ("completed", "Completed"),
            ("blocked", "Blocked"),
        ],
        default="available",
        required=True,
        index=True,
        copy=False,
        tracking=True,
    )
    next_eligible_at = fields.Datetime(index=True, copy=False)
    last_call_at = fields.Datetime(index=True, copy=False, readonly=True)
    cross_campaign_duplicate = fields.Boolean(
        default=False,
        index=True,
        copy=False,
        readonly=True,
    )

    @api.model_create_multi
    def create(self, values_list):
        prepared = []
        operational_membership = False
        if _is_operational(self.env.user) and not _is_global_admin(self.env.user):
            operational_membership = self.env.user._cc_resolve_operational_membership()
        for original in values_list:
            values = dict(original)
            if operational_membership:
                campaign = operational_membership.campaign_id
                supplied_campaign_id = values.get("campaign_id")
                if supplied_campaign_id and supplied_campaign_id != campaign.id:
                    raise AccessError(
                        _("The authenticated campaign membership determines lead scope.")
                    )
                values.update(
                    {
                        "campaign_id": campaign.id,
                        "cc_contact_center_record": True,
                        "cc_source_list_key": values.get("cc_source_list_key")
                        or f"manual:odoo:user:{self.env.user.id}",
                    }
                )
                if operational_membership.role in ("agent", "senior_agent"):
                    supplied_user_id = values.get("user_id")
                    if supplied_user_id and supplied_user_id != self.env.user.id:
                        raise AccessError(_("Agents may create leads only for themselves."))
                    values["user_id"] = self.env.user.id
                    values["queue_state"] = "assigned"
                elif values.get("user_id"):
                    values["queue_state"] = "assigned"
                else:
                    values["queue_state"] = "available"
            if values.get("source_batch_id"):
                values["queue_state"] = "available"
                values["user_id"] = False
                values["cc_contact_center_record"] = True
            prepared.append(values)

        records = super().create(prepared)
        for record in records.filtered(
            lambda lead: lead.cc_contact_center_record
            and lead.user_id
            and lead.queue_state == "assigned"
        ):
            if operational_membership and record.user_id == self.env.user:
                self.env["callcenter.lead.assignment"].sudo().create(
                    {
                        "lead_id": record.id,
                        "campaign_id": record.campaign_id.id,
                        "agent_id": record.user_id.id,
                        "assignment_source": "manual_agent_creation",
                        "assigned_by_id": self.env.user.id,
                        "reason": _("Lead created by the assigned agent."),
                    }
                )
        return records

    def write(self, values):
        governed = self.filtered("cc_contact_center_record")
        if governed and not _is_global_admin(self.env.user):
            queue_capability = self.env.context.get("_cc_lead_queue_capability")
            if queue_capability is QUEUE_SYSTEM_CAPABILITY:
                forbidden = set(values) - QUEUE_CONTROLLED_FIELDS
                if forbidden:
                    raise AccessError(
                        _("The queue service may update only controlled queue fields.")
                    )
            elif (
                self.env.context.get("_cc_crm_transition_capability")
                is CRM_TRANSITION_CAPABILITY
            ):
                forbidden = set(values) - TRANSITION_CONTROLLED_FIELDS
                if forbidden:
                    raise AccessError(
                        _("The transition service may update only governed status fields.")
                    )
            else:
                raise AccessError(
                    _(
                        "Saved call-center leads are locked. Only the Call Center "
                        "Super User or an authorized system process may change them."
                    )
                )
        return super().write(values)

    @api.constrains(
        "campaign_id",
        "normalized_phone",
        "normalized_email",
        "source_external_id",
        "active",
        "cc_contact_center_record",
    )
    def _check_same_campaign_duplicate_keys(self):
        for lead in self:
            if not lead.cc_contact_center_record or not lead.active or not lead.campaign_id:
                continue
            duplicate = self._queue_find_duplicate(
                lead.campaign_id.id,
                _phone_key(lead.normalized_phone or lead.phone_sanitized or lead.phone),
                _email_key(lead.normalized_email or lead.email_from),
                lead.source_external_id,
                exclude_id=lead.id,
            )
            if duplicate:
                raise ValidationError(
                    _("This campaign already contains a lead with the same duplicate key.")
                )

    @api.model
    def _queue_find_duplicate(
        self,
        campaign_id,
        normalized_phone=False,
        normalized_email=False,
        source_external_id=False,
        exclude_id=False,
    ):
        params = [campaign_id]
        exclude_sql = ""
        if exclude_id:
            exclude_sql = " AND id != %s"
            params.append(exclude_id)

        if source_external_id:
            self.env.cr.execute(
                f"""
                    SELECT id
                      FROM crm_lead
                     WHERE campaign_id = %s
                       AND cc_contact_center_record IS TRUE
                       AND active IS TRUE
                       {exclude_sql}
                       AND source_external_id = %s
                     LIMIT 1
                """,
                params + [str(source_external_id).strip()],
            )
            row = self.env.cr.fetchone()
            if row:
                return row[0]

        if normalized_phone:
            self.env.cr.execute(
                f"""
                    SELECT id
                      FROM crm_lead
                     WHERE campaign_id = %s
                       AND cc_contact_center_record IS TRUE
                       AND active IS TRUE
                       {exclude_sql}
                       AND regexp_replace(
                           COALESCE(normalized_phone, phone_sanitized, phone, ''),
                           '[^0-9]', '', 'g'
                       ) = %s
                     LIMIT 1
                """,
                params + [normalized_phone],
            )
            row = self.env.cr.fetchone()
            if row:
                return row[0]
        elif normalized_email:
            self.env.cr.execute(
                f"""
                    SELECT id
                      FROM crm_lead
                     WHERE campaign_id = %s
                       AND cc_contact_center_record IS TRUE
                       AND active IS TRUE
                       {exclude_sql}
                       AND lower(COALESCE(normalized_email, email_from, '')) = %s
                     LIMIT 1
                """,
                params + [normalized_email],
            )
            row = self.env.cr.fetchone()
            if row:
                return row[0]
        return False

    @api.model
    def _queue_has_cross_campaign_duplicate(
        self, campaign_id, normalized_phone=False, normalized_email=False
    ):
        if normalized_phone:
            self.env.cr.execute(
                """
                    SELECT 1
                      FROM crm_lead
                     WHERE campaign_id IS NOT NULL
                       AND campaign_id != %s
                       AND cc_contact_center_record IS TRUE
                       AND active IS TRUE
                       AND regexp_replace(
                           COALESCE(normalized_phone, phone_sanitized, phone, ''),
                           '[^0-9]', '', 'g'
                       ) = %s
                     LIMIT 1
                """,
                [campaign_id, normalized_phone],
            )
        elif normalized_email:
            self.env.cr.execute(
                """
                    SELECT 1
                      FROM crm_lead
                     WHERE campaign_id IS NOT NULL
                       AND campaign_id != %s
                       AND cc_contact_center_record IS TRUE
                       AND active IS TRUE
                       AND lower(COALESCE(normalized_email, email_from, '')) = %s
                     LIMIT 1
                """,
                [campaign_id, normalized_email],
            )
        else:
            return False
        return bool(self.env.cr.fetchone())

    @api.model
    def action_get_next_lead(self):
        membership = self.env.user._cc_resolve_operational_membership()
        if membership.role not in ("agent", "senior_agent"):
            raise AccessError(_("Only campaign agents may request the next lead."))

        campaign = membership.campaign_id
        if not campaign.active or campaign.lifecycle_state not in QUEUE_ACTIVE_LIFECYCLE_STATES:
            raise UserError(_("The active campaign is not eligible for lead distribution."))

        current = self.sudo().search(
            [
                ("cc_contact_center_record", "=", True),
                ("campaign_id", "=", campaign.id),
                ("user_id", "=", self.env.user.id),
                ("queue_state", "=", "assigned"),
                ("active", "=", True),
            ],
            limit=1,
        )
        if current:
            raise UserError(
                _("Disposition the current working lead before requesting another.")
            )

        now = fields.Datetime.now()
        self.env.cr.execute(
            """
                SELECT id
                  FROM crm_lead
                 WHERE campaign_id = %s
                   AND cc_contact_center_record IS TRUE
                   AND active IS TRUE
                   AND user_id IS NULL
                   AND queue_state = 'available'
                   AND (next_eligible_at IS NULL OR next_eligible_at <= %s)
                 ORDER BY priority DESC, create_date ASC, id ASC
                 FOR UPDATE SKIP LOCKED
                 LIMIT 1
            """,
            [campaign.id, now],
        )
        row = self.env.cr.fetchone()
        if not row:
            return {
                "type": "ir.actions.client",
                "tag": "display_notification",
                "params": {
                    "title": _("Lead Queue"),
                    "message": _("NO LEADS CURRENTLY AVAILABLE"),
                    "type": "warning",
                    "sticky": False,
                },
            }

        lead = self.sudo().browse(row[0]).exists()
        lead._callcenter_system_assign(
            self.env.user,
            assignment_source="queue",
            assigned_by=self.env.user,
            reason=_("Assigned by Get Next Lead."),
        )
        return {
            "type": "ir.actions.act_window",
            "name": _("Working Lead"),
            "res_model": "crm.lead",
            "res_id": lead.id,
            "view_mode": "form",
            "target": "current",
        }

    def _callcenter_system_assign(
        self,
        agent,
        assignment_source="queue",
        assigned_by=False,
        reason=False,
    ):
        if not self.env.su:
            raise AccessError(_("Queue assignment is an internal system operation."))
        for lead in self:
            if not lead.cc_contact_center_record or not lead.campaign_id:
                raise ValidationError(_("Only governed campaign leads may be assigned."))
            membership = self.env["cc.campaign.membership"].sudo().search(
                [
                    ("user_id", "=", agent.id),
                    ("campaign_id", "=", lead.campaign_id.id),
                    ("state", "=", "active"),
                    ("role", "in", ("agent", "senior_agent")),
                ],
                limit=1,
            )
            if not membership:
                raise ValidationError(_("The target agent is not active in this campaign."))
            if lead.queue_state != "available" or lead.user_id:
                raise ValidationError(_("The lead is no longer available for assignment."))
            actor = assigned_by or agent
            lead.with_context(
                _cc_lead_queue_capability=QUEUE_SYSTEM_CAPABILITY
            ).write({"user_id": agent.id, "queue_state": "assigned"})
            self.env["callcenter.lead.assignment"].sudo().create(
                {
                    "lead_id": lead.id,
                    "campaign_id": lead.campaign_id.id,
                    "agent_id": agent.id,
                    "assignment_source": assignment_source,
                    "assigned_by_id": actor.id,
                    "reason": reason,
                }
            )
        return True

    def _callcenter_record_attempt(self, occurred_at=False):
        if not self.env.su:
            raise AccessError(_("Call-attempt updates are internal system operations."))
        occurred_at = occurred_at or fields.Datetime.now()
        for lead in self:
            lead.with_context(
                _cc_lead_queue_capability=QUEUE_SYSTEM_CAPABILITY
            ).write(
                {
                    "call_attempt_count": lead.call_attempt_count + 1,
                    "last_call_at": occurred_at,
                }
            )
        return True

    def _callcenter_apply_disposition(
        self, queue_state, next_eligible_at=False, reason=False
    ):
        if not self.env.su:
            raise AccessError(_("Queue disposition is an internal system operation."))
        if queue_state not in {"available", "callback", "completed", "blocked"}:
            raise ValidationError(_("Invalid post-disposition queue state."))
        for lead in self:
            values = {
                "queue_state": queue_state,
                "next_eligible_at": next_eligible_at or False,
            }
            if queue_state == "available":
                values["user_id"] = False
            lead.with_context(
                _cc_lead_queue_capability=QUEUE_SYSTEM_CAPABILITY
            ).write(values)
            self.env["callcenter.lead.assignment"].sudo()._close_for_lead(
                lead, reason=reason
            )
        return True

    def action_callcenter_release_to_queue(self, reason=False):
        if not _is_global_admin(self.env.user):
            raise AccessError(_("Only the Call Center Super User may release leads."))
        for lead in self:
            lead.sudo().with_context(
                _cc_lead_queue_capability=QUEUE_SYSTEM_CAPABILITY
            ).write(
                {
                    "user_id": False,
                    "queue_state": "available",
                    "next_eligible_at": False,
                }
            )
            self.env["callcenter.lead.assignment"].sudo()._close_for_lead(
                lead, reason=reason or _("Released by Call Center Super User.")
            )
        return True


class CallcenterLeadImportBatch(models.Model):
    _name = "callcenter.lead.import.batch"
    _description = "Call Center Lead Import Batch"
    _order = "create_date desc, id desc"

    name = fields.Char(required=True, default=lambda self: _("New Lead Import"))
    campaign_id = fields.Many2one(
        "cc.campaign", required=True, ondelete="restrict", index=True
    )
    upload_file = fields.Binary(required=True, attachment=False)
    filename = fields.Char(required=True)
    file_hash = fields.Char(size=64, readonly=True, index=True, copy=False)
    override_duplicate_file = fields.Boolean(copy=False)
    imported_by_id = fields.Many2one(
        "res.users", readonly=True, copy=False, default=lambda self: self.env.user
    )
    imported_at = fields.Datetime(readonly=True, copy=False)
    total_rows = fields.Integer(readonly=True, copy=False)
    created_count = fields.Integer(readonly=True, copy=False)
    duplicate_count = fields.Integer(readonly=True, copy=False)
    error_count = fields.Integer(readonly=True, copy=False)
    state = fields.Selection(
        [
            ("draft", "Draft"),
            ("validated", "Validated"),
            ("imported", "Imported"),
            ("failed", "Failed"),
        ],
        default="draft",
        required=True,
        readonly=True,
        copy=False,
        index=True,
    )
    line_ids = fields.One2many(
        "callcenter.lead.import.line", "batch_id", readonly=True, copy=False
    )

    @api.model_create_multi
    def create(self, values_list):
        _require_global_admin(self.env)
        return super().create(values_list)

    def write(self, values):
        _require_global_admin(self.env)
        if self.filtered(lambda batch: batch.state != "draft") and {
            "campaign_id",
            "upload_file",
            "filename",
        }.intersection(values):
            raise AccessError(_("Validated import source data is immutable."))
        return super().write(values)

    def unlink(self):
        raise AccessError(_("Lead import batches are retained for audit."))

    def action_validate(self):
        _require_global_admin(self.env)
        for batch in self:
            batch.ensure_one()
            raw = base64.b64decode(batch.upload_file or b"")
            if not raw:
                raise ValidationError(_("The import file is empty."))
            file_hash = hashlib.sha256(raw).hexdigest()
            prior = self.search(
                [
                    ("id", "!=", batch.id),
                    ("campaign_id", "=", batch.campaign_id.id),
                    ("file_hash", "=", file_hash),
                    ("state", "=", "imported"),
                ],
                order="imported_at desc, id desc",
                limit=1,
            )
            if prior and not batch.override_duplicate_file:
                when = fields.Datetime.to_string(prior.imported_at) if prior.imported_at else _("unknown time")
                who = prior.imported_by_id.display_name or _("Administrator")
                raise UserError(
                    _(
                        "This file was already imported on %(when)s by %(who)s. "
                        "Enable the deliberate duplicate-file override to continue.",
                        when=when,
                        who=who,
                    )
                )

            rows = batch._parse_rows(raw)
            batch.line_ids.sudo().unlink()
            seen_phone = set()
            seen_email = set()
            seen_external = set()
            duplicate_count = 0
            error_count = 0
            line_values = []

            for row_number, row in enumerate(rows, start=2):
                mapped = batch._map_row(row)
                name = str(mapped.get("name") or "").strip()
                phone = mapped.get("phone")
                normalized_phone = _phone_key(phone)
                normalized_email = _email_key(mapped.get("email"))
                source_external_id = str(mapped.get("external_id") or "").strip() or False
                status = "valid"
                error = False

                if not name:
                    status, error = "error", _("Name / Contact Name is required.")
                elif not normalized_phone and not normalized_email:
                    status, error = "error", _("A usable phone or fallback email is required.")
                elif (
                    source_external_id and source_external_id in seen_external
                ) or (
                    normalized_phone and normalized_phone in seen_phone
                ) or (
                    not normalized_phone
                    and normalized_email
                    and normalized_email in seen_email
                ):
                    status, error = "duplicate", _("Duplicate row inside this import batch.")
                else:
                    duplicate_id = self.env["crm.lead"].sudo()._queue_find_duplicate(
                        batch.campaign_id.id,
                        normalized_phone,
                        normalized_email,
                        source_external_id,
                    )
                    if duplicate_id:
                        status, error = "duplicate", _("Duplicate lead already exists in this campaign.")

                if status == "valid":
                    if source_external_id:
                        seen_external.add(source_external_id)
                    if normalized_phone:
                        seen_phone.add(normalized_phone)
                    elif normalized_email:
                        seen_email.add(normalized_email)
                elif status == "duplicate":
                    duplicate_count += 1
                else:
                    error_count += 1

                cross_campaign = False
                if status == "valid":
                    cross_campaign = self.env["crm.lead"].sudo()._queue_has_cross_campaign_duplicate(
                        batch.campaign_id.id,
                        normalized_phone,
                        normalized_email,
                    )

                line_values.append(
                    {
                        "batch_id": batch.id,
                        "row_number": row_number,
                        "raw_data": mapped,
                        "normalized_phone": normalized_phone,
                        "normalized_email": normalized_email,
                        "source_external_id": source_external_id,
                        "cross_campaign_duplicate": cross_campaign,
                        "status": status,
                        "error_message": error,
                    }
                )

            if line_values:
                self.env["callcenter.lead.import.line"].sudo().create(line_values)
            super(CallcenterLeadImportBatch, batch).write(
                {
                    "file_hash": file_hash,
                    "total_rows": len(rows),
                    "created_count": 0,
                    "duplicate_count": duplicate_count,
                    "error_count": error_count,
                    "state": "validated",
                }
            )
        return True

    def action_import(self):
        _require_global_admin(self.env)
        for batch in self:
            if batch.state != "validated":
                raise ValidationError(_("Validate the import batch before confirming import."))

            self.env.cr.execute(
                "SELECT id FROM cc_campaign WHERE id = %s FOR UPDATE",
                [batch.campaign_id.id],
            )
            created_count = 0
            for line in batch.line_ids.filtered(lambda item: item.status == "valid"):
                row = line.raw_data or {}
                duplicate_id = self.env["crm.lead"].sudo()._queue_find_duplicate(
                    batch.campaign_id.id,
                    line.normalized_phone,
                    line.normalized_email,
                    line.source_external_id,
                )
                if duplicate_id:
                    line.sudo().write(
                        {
                            "status": "duplicate",
                            "error_message": _("A matching lead was created before import confirmation."),
                        }
                    )
                    continue

                values = batch._lead_values_from_line(line)
                lead = self.env["crm.lead"].with_user(self.env.user).create(values)
                line.sudo().write({"status": "imported", "lead_id": lead.id})
                created_count += 1

            duplicate_count = len(batch.line_ids.filtered(lambda item: item.status == "duplicate"))
            error_count = len(batch.line_ids.filtered(lambda item: item.status == "error"))
            super(CallcenterLeadImportBatch, batch).write(
                {
                    "created_count": created_count,
                    "duplicate_count": duplicate_count,
                    "error_count": error_count,
                    "imported_by_id": self.env.user.id,
                    "imported_at": fields.Datetime.now(),
                    "state": "imported",
                }
            )
        return True

    def _parse_rows(self, raw):
        self.ensure_one()
        filename = (self.filename or "").strip().lower()
        if filename.endswith(".csv"):
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError as exc:
                raise ValidationError(_("CSV files must use UTF-8 encoding.")) from exc
            reader = csv.DictReader(io.StringIO(text))
            if not reader.fieldnames:
                raise ValidationError(_("The CSV file has no header row."))
            return [dict(row) for row in reader]

        if filename.endswith(".xlsx"):
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise UserError(
                    _("XLSX import requires the Python openpyxl package.")
                ) from exc
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            try:
                headers = [str(value or "") for value in next(iterator)]
            except StopIteration as exc:
                raise ValidationError(_("The XLSX file is empty.")) from exc
            return [
                {headers[index]: value for index, value in enumerate(row) if index < len(headers)}
                for row in iterator
            ]

        raise ValidationError(_("Only CSV and XLSX lead imports are supported."))

    def _map_row(self, row):
        mapped = {}
        for key, value in (row or {}).items():
            canonical = _canonical_column(key)
            if canonical in {
                "name",
                "phone",
                "external_id",
                "email",
                "company",
                "priority",
                "source",
                "notes",
            }:
                mapped[canonical] = _json_value(value)
        return mapped

    def _lead_values_from_line(self, line):
        self.ensure_one()
        row = line.raw_data or {}
        values = {
            "name": str(row.get("name") or "").strip(),
            "contact_name": str(row.get("name") or "").strip(),
            "phone": str(row.get("phone") or "").strip() or False,
            "email_from": str(row.get("email") or "").strip() or False,
            "partner_name": str(row.get("company") or "").strip() or False,
            "priority": _priority_value(row.get("priority")),
            "description": str(row.get("notes") or "").strip() or False,
            "campaign_id": self.campaign_id.id,
            "cc_contact_center_record": True,
            "cc_source_list_key": f"import:{self.id}",
            "source_batch_id": self.id,
            "source_external_id": line.source_external_id,
            "queue_state": "available",
            "user_id": False,
            "cross_campaign_duplicate": line.cross_campaign_duplicate,
        }
        source_name = str(row.get("source") or "").strip()
        if source_name:
            Source = self.env["utm.source"].sudo()
            source = Source.search([("name", "=", source_name)], limit=1)
            if not source:
                source = Source.create({"name": source_name})
            values["source_id"] = source.id
        team = self.env["crm.team"].sudo().search(
            [("default_campaign_id", "=", self.campaign_id.legacy_campaign_id.id)],
            order="id",
            limit=1,
        )
        if team:
            values["team_id"] = team.id
        return values


class CallcenterLeadImportLine(models.Model):
    _name = "callcenter.lead.import.line"
    _description = "Call Center Lead Import Line"
    _order = "batch_id, row_number, id"

    batch_id = fields.Many2one(
        "callcenter.lead.import.batch", required=True, ondelete="cascade", index=True
    )
    row_number = fields.Integer(required=True)
    raw_data = fields.Json(required=True, default=dict)
    normalized_phone = fields.Char(index=True)
    normalized_email = fields.Char(index=True)
    source_external_id = fields.Char(index=True)
    cross_campaign_duplicate = fields.Boolean(default=False)
    status = fields.Selection(
        [
            ("valid", "Valid"),
            ("duplicate", "Duplicate"),
            ("error", "Error"),
            ("imported", "Imported"),
        ],
        required=True,
        default="valid",
        index=True,
    )
    lead_id = fields.Many2one("crm.lead", ondelete="restrict", readonly=True, copy=False)
    error_message = fields.Text(readonly=True, copy=False)

    def unlink(self):
        if any(line.batch_id.state == "imported" for line in self):
            raise AccessError(_("Imported lead audit lines cannot be deleted."))
        return super().unlink()


class CallcenterLeadAssignment(models.Model):
    _name = "callcenter.lead.assignment"
    _description = "Call Center Lead Assignment History"
    _order = "assigned_at desc, id desc"

    lead_id = fields.Many2one(
        "crm.lead", required=True, ondelete="restrict", index=True, readonly=True
    )
    campaign_id = fields.Many2one(
        "cc.campaign", required=True, ondelete="restrict", index=True, readonly=True
    )
    agent_id = fields.Many2one(
        "res.users", required=True, ondelete="restrict", index=True, readonly=True
    )
    assigned_at = fields.Datetime(
        required=True, default=fields.Datetime.now, index=True, readonly=True
    )
    released_at = fields.Datetime(index=True, readonly=True)
    assignment_source = fields.Selection(
        [
            ("queue", "Queue"),
            ("admin", "Administrator"),
            ("manual_agent_creation", "Manual Agent Creation"),
            ("transfer", "Transfer"),
        ],
        required=True,
        index=True,
        readonly=True,
    )
    assigned_by_id = fields.Many2one(
        "res.users", required=True, ondelete="restrict", readonly=True
    )
    reason = fields.Text(readonly=True)

    def write(self, values):
        if self.env.context.get(
            "_cc_assignment_history_capability"
        ) is not ASSIGNMENT_HISTORY_CAPABILITY:
            raise AccessError(_("Lead assignment history is append-only."))
        allowed = {"released_at", "reason"}
        if set(values) - allowed:
            raise AccessError(_("Only assignment release metadata may be finalized."))
        if "released_at" in values and any(record.released_at for record in self):
            raise AccessError(_("An assignment release timestamp cannot be overwritten."))
        return super().write(values)

    def unlink(self):
        raise AccessError(_("Lead assignment history is retained for audit."))

    @api.model
    def _close_for_lead(self, lead, reason=False):
        assignment = self.search(
            [
                ("lead_id", "=", lead.id),
                ("released_at", "=", False),
            ],
            order="assigned_at desc, id desc",
            limit=1,
        )
        if assignment:
            values = {"released_at": fields.Datetime.now()}
            if reason:
                values["reason"] = reason
            assignment.with_context(
                _cc_assignment_history_capability=ASSIGNMENT_HISTORY_CAPABILITY
            ).write(values)
        return True
